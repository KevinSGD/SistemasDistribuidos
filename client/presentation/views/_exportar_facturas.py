"""Exportacion a Excel de la pantalla de facturas (filas visibles + hoja de resumen).

Funcion pura sobre datos ya descargados: corre en un hilo de trabajo, sin tocar widgets.
Las facturas son de solo lectura; este archivo es una copia de consulta/auditoria."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from client.infrastructure.factura_pdf import METODOS

from ._exportar import carpeta_reportes

COLUMNAS = ["Número", "Fecha y hora", "Cliente", "Documento", "Método de pago", "Cuenta",
            "Subtotal", "Ahorro", "IVA", "Total"]
HOJA_FACTURAS = "Facturas"
HOJA_RESUMEN = "Resumen"


def nombre_reporte(ahora: datetime | None = None) -> str:
    return f"facturas_{(ahora or datetime.now()).strftime('%Y%m%d_%H%M%S')}.xlsx"


def etiqueta_metodo(metodo: str) -> str:
    return METODOS.get(metodo, str(metodo))


def exportar_facturas_excel(facturas: list, resumen, moneda: str, negocio: str,
                            filtros: str = "", carpeta: Path | None = None,
                            ahora: datetime | None = None) -> Path:
    """`facturas`: FacturaOut en el orden en que se ven en la tabla. `resumen`: ResumenFacturasOut
    del mismo filtro (puede cubrir mas facturas que las exportadas)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ahora = ahora or datetime.now()
    formato_monto = "#,##0"          # el sistema trabaja con pesos enteros
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA_FACTURAS
    ws.append([f"Facturas — {negocio}"])
    ws["A1"].font = Font(bold=True, size=14, color="1E40AF")
    ws.append([f"Generado: {ahora.strftime('%Y-%m-%d %H:%M:%S')} · {len(facturas)} facturas"
               + (f" · Filtros: {filtros}" if filtros else "")])
    ws["A2"].font = Font(color="6B7280")
    ws.append([])
    ws.append(COLUMNAS)
    for celda in ws[4]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1E40AF")
        celda.alignment = Alignment(horizontal="center", vertical="center")
    borde = Border(bottom=Side(style="thin", color="E5E7EB"))
    for f in facturas:
        ws.append([f.numero, f.fecha, f.cliente.nombre, f.cliente.documento,
                   etiqueta_metodo(f.metodo_pago), f.usuario_nombre or "",
                   f.subtotal, f.ahorro, f.iva, f.total])
        fila = ws.max_row
        ws.cell(fila, 2).number_format = "yyyy-mm-dd hh:mm"
        for col in range(7, 11):
            ws.cell(fila, col).number_format = formato_monto
        for celda in ws[fila]:
            celda.border = borde
    for i, ancho in enumerate([14, 18, 28, 16, 20, 22, 14, 14, 14, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A5"
    if facturas:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(COLUMNAS))}{ws.max_row}"

    rs = wb.create_sheet(HOJA_RESUMEN)
    rs.append([f"Resumen de facturas — {negocio}"])
    rs["A1"].font = Font(bold=True, size=14, color="1E40AF")
    rs.append([f"Generado: {ahora.strftime('%Y-%m-%d %H:%M:%S')}"])
    rs["A2"].font = Font(color="6B7280")
    rs.append(["Filtros", filtros or "Ninguno (todas las facturas)"])
    rs.append(["Facturas exportadas", len(facturas)])
    rs.append(["Facturas que cumplen el filtro", resumen.cantidad])
    rs.append(["Rango real", (f"{resumen.desde:%Y-%m-%d %H:%M} a {resumen.hasta:%Y-%m-%d %H:%M}"
                              if resumen.desde and resumen.hasta else "—")])
    rs.append([])
    inicio = rs.max_row + 1
    for etiqueta, valor in (("Subtotal (sin descuento)", resumen.subtotal),
                            ("Ahorro otorgado", resumen.ahorro), ("IVA", resumen.iva),
                            ("Total vendido", resumen.total),
                            ("Ticket promedio", resumen.ticket_promedio)):
        rs.append([etiqueta, valor])
        rs.cell(rs.max_row, 2).number_format = formato_monto
    for fila in range(inicio, rs.max_row + 1):
        rs.cell(fila, 1).font = Font(bold=True)
    rs.append([])
    rs.append(["Ventas por método de pago", "Facturas", "Total"])
    cab = rs.max_row
    for m in resumen.por_metodo:
        rs.append([etiqueta_metodo(m.metodo), m.cantidad, m.total])
        rs.cell(rs.max_row, 3).number_format = formato_monto
    rs.append([])
    rs.append(["Top 5 productos (por unidades)", "Unidades", "Total"])
    cab2 = rs.max_row
    for t in resumen.top_productos:
        rs.append([t.nombre, t.unidades, t.total])
        rs.cell(rs.max_row, 3).number_format = formato_monto
    for fila in (cab, cab2):
        for celda in rs[fila]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1E40AF")
    for col, ancho in zip("ABC", (34, 40, 16)):
        rs.column_dimensions[col].width = ancho

    ruta = (carpeta or carpeta_reportes()) / nombre_reporte(ahora)
    wb.save(ruta)
    return ruta

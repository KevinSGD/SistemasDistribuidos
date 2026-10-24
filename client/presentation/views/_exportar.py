"""Generacion de reportes de inventario (xlsx y pdf) en la carpeta de reportes del cliente.

Funciones puras sobre datos ya descargados: se ejecutan en un hilo de trabajo, sin tocar widgets."""
from __future__ import annotations

import os
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from platformdirs import user_documents_dir

from client.infrastructure import config as cfg

from ..styles.theme import COLOR_NIVEL, fmt_moneda

COLUMNAS = ["ID", "Nombre", "Categoría", "Precio base", "Precio con descuento",
            "Stock", "Stock máx.", "Vencimiento", "Riesgo"]
ETIQUETA_NIVEL = {"critico": "Crítico", "alto": "Alto", "medio": "Medio", "normal": "Normal"}


@dataclass(frozen=True)
class FilaInventario:
    id: int
    nombre: str
    categoria: str
    precio_base: float
    precio_final: float
    descuento: float
    stock: int
    stock_maximo: int
    vencimiento: str | None        # 'YYYY-MM-DD' o None
    nivel: str


def carpeta_reportes() -> Path:
    """Carpeta de reportes: `<INVENTARIO_CLIENT_DIR>/reportes` si esta definida (tests/portatil),
    si no `Documentos/InventarioApp/Reportes` (platformdirs)."""
    base = os.environ.get(cfg.ENV_DIR)
    ruta = Path(base) / "reportes" if base else Path(user_documents_dir()) / "InventarioApp" / "Reportes"
    ruta.mkdir(parents=True, exist_ok=True)
    return ruta


def nombre_reporte(extension: str, ahora: datetime | None = None) -> str:
    return f"inventario_{(ahora or datetime.now()).strftime('%Y%m%d_%H%M%S')}.{extension}"


def _venc(f: FilaInventario) -> str:
    return f.vencimiento or "Sin vencimiento"


def exportar_excel(filas: list[FilaInventario], moneda: str, negocio: str,
                   carpeta: Path | None = None, ahora: datetime | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ahora = ahora or datetime.now()
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    ws.append([f"Inventario — {negocio}"])
    ws["A1"].font = Font(bold=True, size=14, color="1E40AF")
    ws.append([f"Generado: {ahora.strftime('%Y-%m-%d %H:%M:%S')} · {len(filas)} productos"])
    ws["A2"].font = Font(color="6B7280")
    ws.append([])
    ws.append(COLUMNAS)
    for celda in ws[4]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1E40AF")
        celda.alignment = Alignment(horizontal="center", vertical="center")
    borde = Border(bottom=Side(style="thin", color="E5E7EB"))
    formato_precio = '#,##0.00' if moneda.upper() != "COP" else '#,##0'
    for f in filas:
        ws.append([f.id, f.nombre, f.categoria, f.precio_base, f.precio_final, f.stock,
                   f.stock_maximo, _venc(f), ETIQUETA_NIVEL.get(f.nivel, "Normal")])
        fondo = COLOR_NIVEL.get(f.nivel, COLOR_NIVEL["normal"])[1].lstrip("#")
        for celda in ws[ws.max_row]:
            celda.fill = PatternFill("solid", fgColor=fondo)
            celda.border = borde
        ws.cell(ws.max_row, 4).number_format = formato_precio
        ws.cell(ws.max_row, 5).number_format = formato_precio
    anchos = [8, 30, 18, 14, 20, 9, 11, 16, 11]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A5"
    if filas:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(COLUMNAS))}{ws.max_row}"
    ruta = (carpeta or carpeta_reportes()) / nombre_reporte("xlsx", ahora)
    wb.save(ruta)
    return ruta


def exportar_pdf(filas: list[FilaInventario], moneda: str, negocio: str,
                 carpeta: Path | None = None, ahora: datetime | None = None) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    ahora = ahora or datetime.now()
    ruta = (carpeta or carpeta_reportes()) / nombre_reporte("pdf", ahora)
    doc = SimpleDocTemplate(str(ruta), pagesize=landscape(letter), leftMargin=1.5 * cm,
                            rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            title=f"Inventario — {negocio}", author="InventarioApp")
    estilos = getSampleStyleSheet()
    celda = ParagraphStyle("celda", parent=estilos["BodyText"], fontSize=8.5, leading=10.5)
    elementos = [
        Paragraph(f"<b>Reporte de inventario — {escape(negocio)}</b>", estilos["Title"]),
        Paragraph(f"Generado: {ahora.strftime('%Y-%m-%d %H:%M:%S')} · {len(filas)} productos",
                  estilos["Normal"]),
        Spacer(1, 14),
    ]
    datos = [COLUMNAS]
    for f in filas:
        datos.append([str(f.id), Paragraph(escape(f.nombre), celda), Paragraph(escape(f.categoria), celda),
                      fmt_moneda(f.precio_base, moneda), fmt_moneda(f.precio_final, moneda),
                      str(f.stock), str(f.stock_maximo), _venc(f),
                      ETIQUETA_NIVEL.get(f.nivel, "Normal")])
    tabla = Table(datos, repeatRows=1, colWidths=[1.2 * cm, 5.6 * cm, 3.6 * cm, 2.8 * cm, 3.4 * cm,
                                                  1.6 * cm, 1.9 * cm, 2.8 * cm, 2 * cm])
    estilo = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E40AF")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (3, 0), (6, -1), "RIGHT"),
        ("ALIGN", (7, 0), (-1, -1), "CENTER"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E7EB")),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ])
    for i, f in enumerate(filas, 1):
        fondo = COLOR_NIVEL.get(f.nivel, COLOR_NIVEL["normal"])[1]
        estilo.add("BACKGROUND", (0, i), (-1, i), colors.HexColor(fondo))
    tabla.setStyle(estilo)
    elementos.append(tabla)

    def pie(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawRightString(landscape(letter)[0] - 1.5 * cm, 0.9 * cm, f"Página {documento.page}")
        canvas.restoreState()

    doc.build(elementos, onFirstPage=pie, onLaterPages=pie)
    return ruta

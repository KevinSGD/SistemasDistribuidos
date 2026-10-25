"""Pantalla "Facturas" del administrador contra un servidor REAL propio (BD desechable)."""
from __future__ import annotations

import sqlite3
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pytest
from PyQt6.QtWidgets import QApplication

from client.application.sesion import Aplicacion, OpcionesCliente
from client.infrastructure.api_client import ApiClient, ApiError, ErrorConexion
from client.presentation.styles.theme import fmt_moneda
from client.presentation.views import _comunes as comunes
from client.presentation.views import facturas_view as fv
from client.presentation.views.facturas_view import FacturasView
from client.presentation.views.main_window import MainWindow, VISTAS
from client.presentation.widgets import Toast
from shared.schemas import ClienteIn, CompraIn, ItemCompraIn, PagoIn, ProductoIn, UsuarioIn
from tests.client.conftest import ADMIN_PW, COMP_PW, Servidor

RAPIDO = OpcionesCliente(throttle_actividad=0.2, tick_ms=100, health_ms=500, sondeo_ms=60000,
                         ping_ms=60000, backoff=(0.1, 0.2), solo_espontaneos=False,
                         cierre_logout_ms=200)
CLAVE = "clave-segura-123"
CLIENTES = [("José Pérez", "1001"), ("María López", "1002"), ("Carlos Ruiz", "1003"),
            ("Ana Gómez", "1004"), ("Luis Núñez", "1005"), ("Sofía Vargas", "1006"),
            ("", ""), ("Pedro Pérez", "1008")]
#: dias hacia atras de cada una de las 15 facturas (las 3 primeras son de hoy)
DIAS = [0, 0, 0, 1, 2, 3, 5, 6, 8, 10, 12, 15, 18, 20, 25]


def sembrar_facturas(srv: Servidor, n: int = 15) -> dict:
    """Crea compradores, productos y `n` facturas por la API; reparte las fechas con SQL."""
    admin = ApiClient(srv.url)
    admin.login("admin", ADMIN_PW)
    for nombre, usuario in (("Ana Compradora", "ana"), ("Beto Comprador", "beto")):
        admin.crear_usuario(UsuarioIn(nombre=nombre, usuario=usuario, password=CLAVE))
    productos = [admin.crear_producto(ProductoIn(
        nombre=nom, categoria="Varios", precio_base=precio, stock=5000, stock_maximo=9000))
        for nom, precio in (("Leche 1 L", 4200), ("Pan tajado", 6800), ("Arroz 500 g", 3100),
                            ("Aceite 1 L", 15900), ("Huevos x12", 12500), ("Café molido", 18700))]
    apis = []
    for usuario, clave in (("comprador", COMP_PW), ("ana", CLAVE), ("beto", CLAVE)):
        a = ApiClient(srv.url)
        a.login(usuario, clave)
        apis.append(a)
    numeros = []
    for i in range(n):
        nombre, doc = CLIENTES[i % len(CLIENTES)]
        metodo = "efectivo" if i % 3 else "tarjeta_simulada"
        items = [ItemCompraIn(producto_id=productos[(i + k) % 6].id, cantidad=1 + (i + k) % 4)
                 for k in range(1 + i % 3)]
        pago = PagoIn(metodo=metodo, efectivo_recibido=500000 if metodo == "efectivo" else None)
        f = apis[i % 3].comprar(CompraIn(items=items, cliente=ClienteIn(nombre=nombre, documento=doc),
                                         pago=pago))
        numeros.append(f.numero)
    for a in apis:
        a.cerrar()
    admin.cerrar()
    hoy = date.today()
    con = sqlite3.connect(srv.dir / "test.db", timeout=10)
    try:
        for i, num in enumerate(numeros):
            dias = DIAS[i % len(DIAS)]
            if dias == 0:
                fecha = datetime.combine(hoy, time(8 + i, 15))
            else:
                fecha = datetime.combine(hoy - timedelta(days=dias), time(9 + i % 8, 30 + i))
            con.execute("UPDATE facturas SET fecha=? WHERE numero=?",
                        (fecha.isoformat(sep=" ", timespec="microseconds"), num))
        con.commit()
    finally:
        con.close()
    return {"numeros": numeros, "productos": productos}


@pytest.fixture(scope="module")
def srv_fact(tmp_path_factory):
    s = Servidor(tmp_path_factory.mktemp("srv_facturas")).iniciar()
    datos = sembrar_facturas(s)
    yield s, datos
    s.detener()


@pytest.fixture
def api(srv_fact):
    c = ApiClient(srv_fact[0].url)
    c.login("admin", ADMIN_PW)
    yield c
    c.cerrar()


@pytest.fixture
def api_comp(srv_fact):
    c = ApiClient(srv_fact[0].url)
    c.login("comprador", COMP_PW)
    yield c
    c.cerrar()


@pytest.fixture
def app_admin(qapp, srv_fact, espera):
    a = Aplicacion(srv_fact[0].url, RAPIDO, timeout=4.0)
    a.iniciar()
    a.login.input_usuario.setText("admin")
    a.login.input_password.setText(ADMIN_PW)
    a.login.btn_login.click()
    assert espera(lambda: a.sesion is not None and isinstance(a.ventana, MainWindow))
    yield a
    if a.sesion is not None:
        a._cerrando = False
        a._cerrar_sesion_local("fin", avisar_servidor=False)
    if a.login is not None:
        a.login.cerrada.disconnect()
        a.login.close()
        a.login.deleteLater()
    a.deleteLater()
    QApplication.processEvents()


@pytest.fixture
def vista(app_admin, espera) -> FacturasView:
    w = app_admin.ventana
    assert espera(lambda: w.vistas["dashboard"].cargada)
    w.botones["facturas"].click()
    v = w.vistas["facturas"]
    assert espera(lambda: v.cargada and v.tabla.rowCount() > 0), "la vista no cargó"
    return v


def num_api(api, **kw):
    return [f.numero for f in api.listar_facturas(limite=1000, **kw)]


def dinero(v, moneda="$"):
    return fmt_moneda(v, moneda)


def esperar_numeros(espera, v, esperados, timeout=8.0):
    return espera(lambda: v.numeros() == esperados and v.lbl_conteo.text() != "Actualizando…", timeout)


# ---------------------------------------------------------------------------------------------
def test_la_pantalla_existe_solo_en_la_ventana_del_administrador():
    assert [k for k, _e, _c in VISTAS].index("facturas") == [k for k, _e, _c in VISTAS].index("alertas") + 1
    from client.presentation.views import buyer_window
    assert not hasattr(buyer_window, "FacturasView")


def test_el_comprador_no_puede_pedir_el_resumen(api_comp):
    with pytest.raises(ApiError) as e:
        api_comp.resumen_facturas()
    assert e.value.http == 403 and e.value.codigo == "ROL_NO_PERMITIDO"
    # y su historial solo trae las suyas
    assert 0 < len(api_comp.listar_facturas()) < 15
    assert all(f.usuario_nombre == "Comprador" for f in api_comp.listar_facturas())


def test_carga_inicial_tabla_y_tarjetas_coinciden_con_el_servidor(vista, api, srv_fact):
    r = api.resumen_facturas()
    assert vista.pagina == fv.TABLA
    assert vista.numeros() == num_api(api)                      # mas recientes primero
    total = len(num_api(api))
    assert vista.t_cantidad.valor == str(total) == str(r.cantidad) and total >= 15
    assert vista.t_total.valor == dinero(r.total)
    assert vista.t_ahorro.valor == dinero(r.ahorro)
    assert vista.t_iva.valor == dinero(r.iva)
    assert vista.t_ticket.valor == dinero(r.ticket_promedio)
    assert vista.lbl_conteo.text() == f"{total} facturas"
    top = vista.panel_top.textos()
    assert len(top) == min(5, len(r.top_productos))
    assert top[0][0].endswith(r.top_productos[0].nombre) and top[0][2] == dinero(r.top_productos[0].total)
    assert sum(int(q.split()[0]) for _n, q, _t in vista.panel_metodos.textos()) == total
    # columnas: cuenta (no el cliente), ahorro y total alineados a la derecha con formato de dinero
    fila = vista.tabla.item(0, fv.COL_TOTAL)
    assert fila.text().startswith("$") and int(fila.textAlignment()) & 0x2   # AlignRight
    f0 = api.listar_facturas(limite=1)[0]
    assert vista.tabla.texto(0, fv.COL_CUENTA) == f0.usuario_nombre
    assert vista.tabla.texto(0, fv.COL_CLIENTE) == (f0.cliente.nombre or fv.SIN_DATO)


def test_la_tabla_ordena_por_columnas_y_por_valor(vista, api):
    vista.tabla.sortByColumn(fv.COL_TOTAL, fv.Qt.SortOrder.AscendingOrder)
    totales = [vista.tabla.item(r, fv.COL_TOTAL).data(fv.Qt.ItemDataRole.UserRole)
               for r in range(vista.tabla.rowCount())]
    assert totales == sorted(totales)
    vista.tabla.sortByColumn(fv.COL_NUMERO, fv.Qt.SortOrder.AscendingOrder)
    assert vista.numeros() == sorted(vista.numeros())
    vista.tabla.sortByColumn(fv.COL_FECHA, fv.Qt.SortOrder.DescendingOrder)
    assert vista.numeros() == num_api(api)


def test_filtro_por_fecha_con_atajos(vista, api, espera):
    hoy = datetime.combine(date.today(), time.min)
    vista.aplicar_atajo("hoy")
    esperados = num_api(api, desde=hoy, hasta=datetime.combine(date.today(), time.max))
    assert len(esperados) >= 3
    assert esperar_numeros(espera, vista, esperados)
    r = api.resumen_facturas(desde=hoy, hasta=datetime.combine(date.today(), time.max))
    assert vista.t_cantidad.valor == str(len(esperados)) and vista.t_total.valor == dinero(r.total)
    assert vista.atajos["hoy"].isChecked() and not vista.atajos["todo"].isChecked()

    vista.aplicar_atajo("7dias")
    desde7 = datetime.combine(date.today() - timedelta(days=6), time.min)
    esperados = num_api(api, desde=desde7, hasta=datetime.combine(date.today(), time.max))
    assert 3 < len(esperados) < len(num_api(api))
    assert esperar_numeros(espera, vista, esperados)

    vista.aplicar_atajo("mes")
    desde_mes = datetime.combine(date.today().replace(day=1), time.min)
    esperados = num_api(api, desde=desde_mes, hasta=datetime.combine(date.today(), time.max))
    assert esperar_numeros(espera, vista, esperados)

    vista.aplicar_atajo("todo")
    assert esperar_numeros(espera, vista, num_api(api))
    assert vista.t_cantidad.valor == str(len(num_api(api)))


def test_filtro_por_fechas_manuales_desmarca_el_atajo(vista, api, espera):
    vista.aplicar_atajo("hoy")
    assert espera(lambda: vista.t_cantidad.valor == str(len(num_api(api, desde=datetime.combine(date.today(), time.min)))))
    ayer = date.today() - timedelta(days=1)
    vista.f_desde.setDate(fv.QDate(ayer.year, ayer.month, ayer.day))
    vista.f_hasta.setDate(fv.QDate(ayer.year, ayer.month, ayer.day))
    assert not any(b.isChecked() for b in vista.atajos.values())
    esperados = num_api(api, desde=datetime.combine(ayer, time.min),
                        hasta=datetime.combine(ayer, time.max))
    assert len(esperados) == 1
    assert esperar_numeros(espera, vista, esperados)


def test_busqueda_por_texto_con_retardo_sin_tildes(vista, api, espera):
    vista.busqueda.setText("perez")                       # «José Pérez» y «Pedro Pérez»
    assert vista._timer_busqueda.isActive()               # aun no consulta: hay retardo
    esperados = num_api(api, q="perez")
    assert len(esperados) >= 2
    assert esperar_numeros(espera, vista, esperados)
    r = api.resumen_facturas(q="perez")
    assert vista.t_cantidad.valor == str(r.cantidad) and vista.t_total.valor == dinero(r.total)
    # por numero de factura y por documento
    vista.busqueda.setText(esperados[0])
    assert esperar_numeros(espera, vista, [esperados[0]])
    vista.busqueda.setText("1005")
    assert esperar_numeros(espera, vista, num_api(api, q="1005"))
    assert all(vista.tabla.texto(r, fv.COL_DOC) == "1005" for r in range(vista.tabla.rowCount()))


def test_filtro_por_metodo_de_pago(vista, api, espera):
    i = vista.combo_metodo.findData("tarjeta_simulada")
    vista.combo_metodo.setCurrentIndex(i)
    esperados = num_api(api, metodo_pago="tarjeta_simulada")
    assert 0 < len(esperados) < len(num_api(api))
    assert esperar_numeros(espera, vista, esperados)
    assert {vista.tabla.texto(r, fv.COL_METODO) for r in range(vista.tabla.rowCount())} \
        == {"Tarjeta (simulada)"}
    r = api.resumen_facturas(metodo_pago="tarjeta_simulada")
    assert vista.t_total.valor == dinero(r.total)
    assert [m[0] for m in vista.panel_metodos.textos()] == ["Tarjeta (simulada)"]


def test_filtros_combinados_sin_resultados_y_limpiar(vista, api, espera):
    vista.busqueda.setText("zzz-no-existe")
    assert espera(lambda: vista.pagina == fv.SIN_RESULTADOS)
    assert vista.t_cantidad.valor == "0" and vista.t_total.valor == dinero(0)
    assert not vista.btn_excel.isEnabled()
    vista.estado_sin_resultados.accion.emit()              # "Limpiar filtros"
    assert espera(lambda: vista.pagina == fv.TABLA and vista.tabla.rowCount() == len(num_api(api)))
    assert vista.busqueda.text() == "" and vista.combo_metodo.currentIndex() == 0
    assert not vista.hay_filtros()


def test_ver_factura_con_boton_y_doble_clic(vista, api, espera):
    assert not vista.btn_ver.isEnabled()
    vista.tabla.selectRow(2)
    assert vista.btn_ver.isEnabled()
    numero = vista.tabla.texto(2, fv.COL_NUMERO)
    vista.btn_ver.click()
    d = vista.dialogo_factura
    assert d is not None and d.datos["numero"] == numero
    assert d.btn_pdf.text() == "Guardar / Imprimir PDF"
    assert "Factura emitida" in d.lbl_estado.text() and "Cuenta:" in d.lbl_estado.text()
    d.close()
    assert espera(lambda: vista.dialogo_factura is None)
    vista.tabla.clearSelection()
    vista.tabla.setCurrentCell(-1, -1)
    vista.tabla.selectRow(0)
    vista.tabla.itemDoubleClicked.emit(vista.tabla.item(0, 0))
    assert vista.dialogo_factura.datos["numero"] == vista.tabla.texto(0, fv.COL_NUMERO)
    vista.dialogo_factura.close()


def test_exportar_excel_filas_visibles_y_hoja_de_resumen(vista, api, espera, monkeypatch, tmp_path):
    from openpyxl import load_workbook
    vista.combo_metodo.setCurrentIndex(vista.combo_metodo.findData("efectivo"))
    esperados = num_api(api, metodo_pago="efectivo")
    assert esperar_numeros(espera, vista, esperados)
    vista.btn_excel.click()
    assert vista.btn_excel.text() == "Exportando…" and not vista.btn_excel.isEnabled()
    assert espera(lambda: vista._ultimo_reporte is not None)
    ruta = vista._ultimo_reporte
    assert ruta.exists() and ruta.suffix == ".xlsx" and ruta.name.startswith("facturas_")
    assert ruta.parent == Path(__import__("os").environ["INVENTARIO_CLIENT_DIR"]) / "reportes"
    assert vista.btn_excel.text() == "Exportar a Excel" and vista.btn_excel.isEnabled()
    assert not vista.barra_reporte.isHidden() and ruta.name in vista.lbl_reporte.text()
    assert Toast.activos(vista.window())

    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Facturas", "Resumen"]
    ws = wb["Facturas"]
    filas = [[c.value for c in fila] for fila in ws.iter_rows(min_row=5)]
    assert [f[0] for f in filas] == esperados
    por_num = {f.numero: f for f in api.listar_facturas(limite=1000)}
    for f in filas:
        o = por_num[f[0]]
        assert f[4] == "Efectivo" and f[5] == o.usuario_nombre and f[9] == o.total
        assert (f[2] or "") == o.cliente.nombre and isinstance(f[1], datetime)
    rs = api.resumen_facturas(metodo_pago="efectivo")
    celdas = {r[0].value: r[1].value for r in wb["Resumen"].iter_rows() if r[0].value}
    assert celdas["Total vendido"] == rs.total and celdas["Ahorro otorgado"] == rs.ahorro
    assert celdas["Facturas exportadas"] == len(esperados) == celdas["Facturas que cumplen el filtro"]
    assert "método Efectivo" in celdas["Filtros"]

    abiertos = []
    monkeypatch.setattr(comunes, "abrir_ruta", lambda r: abiertos.append(Path(r)))
    vista.btn_abrir_archivo.click()
    vista.btn_abrir_carpeta.click()
    assert abiertos == [ruta, ruta.parent]


def test_exportar_sin_filas_avisa_y_no_crea_archivo(vista, espera):
    vista.busqueda.setText("zzz-no-existe")
    assert espera(lambda: vista.pagina == fv.SIN_RESULTADOS)
    vista.exportar()
    assert vista._ultimo_reporte is None
    assert any("No hay facturas" in t.text() for t in Toast.activos(vista.window()))


def test_evento_factura_creada_refresca_sin_pisar_filtros_ni_contar_actividad(
        vista, app_admin, api, espera, monkeypatch, srv_fact):
    marcas, registros = [], []
    monkeypatch.setattr(app_admin, "marcar_actividad_vista", lambda: marcas.append(1))
    original = app_admin.api.registrar_actividad
    monkeypatch.setattr(app_admin.api, "registrar_actividad",
                        lambda: (registros.append(1), original())[1])
    vista.combo_metodo.setCurrentIndex(vista.combo_metodo.findData("tarjeta_simulada"))
    antes = num_api(api, metodo_pago="tarjeta_simulada")
    assert esperar_numeros(espera, vista, antes)
    vista.tabla.selectRow(1)
    seleccionada = vista.tabla.texto(1, fv.COL_NUMERO)

    prod = srv_fact[1]["productos"][0]
    comprador = ApiClient(srv_fact[0].url)
    comprador.login("beto", CLAVE)
    try:
        nueva = comprador.comprar(CompraIn(
            items=[ItemCompraIn(producto_id=prod.id, cantidad=2)],
            cliente=ClienteIn(nombre="Cliente En Vivo", documento="777"),
            pago=PagoIn(metodo="tarjeta_simulada")))
        efectivo = comprador.comprar(CompraIn(
            items=[ItemCompraIn(producto_id=prod.id, cantidad=1)],
            cliente=ClienteIn(nombre="Otro", documento="778"),
            pago=PagoIn(metodo="efectivo", efectivo_recibido=100000)))
    finally:
        comprador.cerrar()
    assert espera(lambda: nueva.numero in vista.numeros())
    # el filtro de tarjeta sigue vigente: la de efectivo no aparece; la seleccion se conserva
    assert efectivo.numero not in vista.numeros()
    assert vista.combo_metodo.currentData() == "tarjeta_simulada"
    assert vista.numeros()[0] == nueva.numero == num_api(api, metodo_pago="tarjeta_simulada")[0]
    assert vista.numero_seleccionado() == seleccionada
    assert vista.t_cantidad.valor == str(len(antes) + 1)
    assert any(f"Nueva factura {nueva.numero}" in t.text() for t in Toast.activos(vista.window()))
    assert marcas == [] and registros == []


def test_el_evento_refresca_aunque_la_vista_no_este_visible(app_admin, api, espera, srv_fact):
    w = app_admin.ventana
    v = w.vistas["facturas"]
    w.botones["facturas"].click()
    assert espera(lambda: v.cargada)
    w.botones["dashboard"].click()
    n = v.tabla.rowCount()
    prod = srv_fact[1]["productos"][1]
    c = ApiClient(srv_fact[0].url)
    c.login("ana", CLAVE)
    try:
        c.comprar(CompraIn(items=[ItemCompraIn(producto_id=prod.id, cantidad=1)],
                           pago=PagoIn(metodo="tarjeta_simulada")))
    finally:
        c.cerrar()
    assert espera(lambda: v.tabla.rowCount() == n + 1)
    # sin aviso emergente cuando no esta en pantalla
    assert not any("Nueva factura" in t.text() for t in Toast.activos(w))


def test_estado_de_error_con_reintentar(vista, app_admin, espera, monkeypatch):
    original = app_admin.api.listar_facturas

    def falla(*a, **k):
        raise ErrorConexion("sin red")
    monkeypatch.setattr(app_admin.api, "listar_facturas", falla)
    vista.btn_actualizar.click()
    assert espera(lambda: vista.pagina == fv.ERROR)
    assert vista.estado_error.boton is not None and vista.t_cantidad.valor == "–"
    assert not vista.btn_excel.isEnabled()
    # un sondeo silencioso fallido no tapa datos buenos, pero aqui ya hay error: se mantiene
    monkeypatch.setattr(app_admin.api, "listar_facturas", original)
    vista.estado_error.boton.click()                        # Reintentar
    assert espera(lambda: vista.pagina == fv.TABLA and vista.tabla.rowCount() > 0)


def test_un_fallo_silencioso_conserva_los_datos_buenos(vista, app_admin, espera, monkeypatch):
    def falla(*a, **k):
        raise ErrorConexion("sin red")
    monkeypatch.setattr(app_admin.api, "listar_facturas", falla)
    n = vista.tabla.rowCount()
    app_admin.sondeo.emit()
    espera(lambda: False, timeout=0.6)
    assert vista.pagina == fv.TABLA and vista.tabla.rowCount() == n


def test_indicador_mostrando_n_de_m(vista, api, espera, monkeypatch):
    monkeypatch.setattr(fv, "LIMITE_UI", 4)
    vista.recargar(manual=True)
    assert espera(lambda: vista.tabla.rowCount() == 4)
    total = len(num_api(api))
    assert vista.lbl_conteo.text().startswith(f"Mostrando 4 de {total} facturas")
    assert vista.t_cantidad.valor == str(total)
    monkeypatch.setattr(fv, "LIMITE_UI", 500)


def test_navegar_a_la_vista_cuenta_como_actividad_y_pide_datos(app_admin, espera, monkeypatch):
    w = app_admin.ventana
    marcas = []
    monkeypatch.setattr(app_admin, "marcar_actividad_vista", lambda: marcas.append(1))
    w.botones["facturas"].click()
    assert marcas == [1]
    assert espera(lambda: w.vistas["facturas"].cargada)
